#!/usr/bin/env python3
"""Render the checked-in integration guide Markdown as a compact DOCX."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "GROK-WEB-INTEGRATION-GUIDE.zh-CN.md"
OUTPUT = ROOT / "GROK-WEB-INTEGRATION-GUIDE.zh-CN.docx"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
MUTED = "667085"


def set_font(run, size=11, bold=False, color="000000", name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shading = props.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        props.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    props = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    props.append(repeat)


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    props = table._tbl.tblPr
    width = props.first_child_found_in("w:tblW")
    width.set(qn("w:w"), "9360")
    width.set(qn("w:type"), "dxa")
    indent = props.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        props.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(value))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tcw = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tcw.set(qn("w:w"), str(widths[idx]))
            tcw.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def proportional_widths(rows):
    cols = len(rows[0])
    weights = []
    for idx in range(cols):
        longest = max(len(re.sub(r"[`*]", "", row[idx])) for row in rows)
        weights.append(max(8, min(longest, 30)))
    total = sum(weights)
    widths = [max(1050, round(9360 * value / total)) for value in weights]
    scale = 9360 / sum(widths)
    widths = [round(value * scale) for value in widths]
    widths[-1] += 9360 - sum(widths)
    return widths


def add_inline(paragraph, text, size=11, color="000000", bold=False):
    text = text.replace("<br>", "")
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        code = part.startswith("`") and part.endswith("`")
        strong = part.startswith("**") and part.endswith("**")
        value = part[1:-1] if code else part[2:-2] if strong else part
        run = paragraph.add_run(value)
        set_font(run, size=size, bold=bold or strong, color=color, name="Consolas" if code else "Calibri")
        if code:
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "F2F4F7")
            run._element.get_or_add_rPr().append(shading)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    add_inline(header, "Grok Web 模型对接与实测指南", size=9, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_inline(footer, "Grok2API v3.1.3-upscale.2 · 2026-08-19", size=8.5, color=MUTED)


def add_title_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(56)
    p.paragraph_format.space_after = Pt(8)
    add_inline(p, "Grok Web 模型对接与实测指南", size=27, color=DARK_BLUE, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(24)
    add_inline(p, "文字 · 文生图 · 图片编辑 · 视频 · 代理与质量守护", size=13, color=MUTED)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(36)
    add_inline(p, "测试日期  2026-08-19    |    部署版本  v3.1.3-upscale.2", size=10.5, color=MUTED, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    add_inline(p, "本文档以真实落盘文件为准。HTTP 200 但没有有效媒体文件不计为成功。", size=11, color=DARK_BLUE, bold=True)


def parse_table(lines):
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    if len(rows) > 1 and all(re.fullmatch(r":?-+:?", cell.replace(" ", "")) for cell in rows[1]):
        rows.pop(1)
    return rows


def build():
    doc = Document()
    style_document(doc)
    add_title_block(doc)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 4  # The title and metadata are represented by the custom cover.
    in_code = False
    code_lines = []
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if stripped.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.18)
                p.paragraph_format.right_indent = Inches(0.18)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                add_inline(p, "\n".join(code_lines), size=8.5, color="233044")
                props = p._p.get_or_add_pPr()
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), LIGHT_FILL)
                props.append(shading)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(raw)
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and lines[index + 1].strip().startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            rows = parse_table(table_lines)
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            set_repeat_header(table.rows[0])
            for ridx, values in enumerate(rows):
                for cidx, value in enumerate(values):
                    cell = table.cell(ridx, cidx)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.08
                    add_inline(paragraph, value, size=8.5, color=DARK_BLUE if ridx == 0 else "000000", bold=ridx == 0)
                    if ridx == 0:
                        set_cell_shading(cell, HEADER_FILL)
            set_table_geometry(table, proportional_widths(rows))
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        if not stripped:
            index += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)", stripped)
        if heading:
            level = len(heading.group(1))
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline(p, heading.group(2), size={1: 16, 2: 13, 3: 12}[level], color=BLUE if level < 3 else DARK_BLUE, bold=True)
            index += 1
            continue
        bullet = re.match(r"^-\s+(.*)", stripped)
        numbered = re.match(r"^\d+\.\s+(.*)", stripped)
        if bullet or numbered:
            p = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            p.paragraph_format.left_indent = Inches(0.375)
            p.paragraph_format.first_line_indent = Inches(-0.188)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.25
            add_inline(p, (bullet or numbered).group(1))
            index += 1
            continue
        p = doc.add_paragraph()
        add_inline(p, stripped.replace("  ", " "))
        index += 1
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
